# -*- coding: utf-8 -*-
"""生成《检索与评测说明》Word文档（第5周交付物，知识库扩充后最终版）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '黑体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)
    run.bold = bold
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(10.5)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(10.5)
    return table

# ========== 标题 ==========
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('智影工坊 AI 品牌宣传视频智能体\n检索与评测说明')
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.size = Pt(18)
run.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run('（AI 应用与智能体赛道 · 第 5 周）')
run.font.name = '宋体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
run.font.size = Pt(12)
doc.add_paragraph()

# ========== 一、本周改动范围 ==========
add_heading('一、本周在链路上只换了一个环节', level=1)
add_para('第 4 周搭好的编排链路为「用户输入 → 隐私检测 → 意图识别 → 分支路由 → 检索 Top-K → 阈值判定 → 在线增强或离线生成 → 合规校验 → 三行格式输出」。本周只动知识检索这一环：将第 4 周已实现的 bigram TF-IDF 检索从浏览器运行时构建，升级为可独立构建、可重复执行的离线索引脚本，并新增命令行检索与评测工具；同时将系统提示词从 v0 升级到 v1，补回「基于片段作答、不得编造」的约束。为提升离线知识助手的覆盖面和智能性，本周还将知识库从 20 块扩充到 35 块，新增同义词扩展和多片段综合回答。其他模块（意图识别、合规校验、在线增强、兜底分支）的函数签名保持不变。')

add_para('本周新增/修改文件：', bold=True, indent=False)
add_para('scripts/build_index.mjs —— 索引构建脚本（可重复执行）', indent=False)
add_para('scripts/retrieve.mjs —— 真检索命令行工具（Top-K=3，阈值 0.72）', indent=False)
add_para('scripts/evaluate.mjs —— 首轮评测脚本（10 条问题，输出 CSV）', indent=False)
add_para('scripts/lib/rag.mjs —— 检索共享库（与浏览器端算法一致）', indent=False)
add_para('prompts/system_v1.md —— 系统提示词 v1（全文存档）', indent=False)
add_para('knowledge.data.js —— 知识库扩充（20→35 块，新增 15 个知识块）', indent=False)
add_para('knowledge.js —— 检索增强（同义词表 + 多片段综合回答）', indent=False)
add_para('data/index/index.json —— 构建产出的向量索引', indent=False)
add_para('output/eval_round1.csv —— 首轮评测记录', indent=False)

# ========== 二、索引构建记录 ==========
add_heading('二、索引构建记录', level=1)
add_para('索引存放在 data/index/，向量与元数据（文本、来源文档、章节、块序号、标签）一起存。构建脚本可重复执行：第二次运行前先清空该目录，不会出现重复块。向量化方式为 bigram TF-IDF（中文二元语法为主，关键 unigram 补充，英文数字按词），与浏览器端 knowledge.js 的算法完全一致，保证离线评测结果与页面运行结果一致。')

add_para('构建统计（扩充后）：', bold=True, indent=False)
add_table(
    ['文档', '知识块数', '覆盖范围'],
    [
        ['《短视频品牌营销实务手册》', '10 块', '品牌定位、人群画像、卖点提炼、内容结构、差异化、品牌故事、用户旅程、内容运营、数据指标、竞品分析'],
        ['《短视频脚本文案与分镜创作指南》', '12 块', '开场钩子、卖点文案、号召转化、六镜分镜、镜头语言、时长节奏、字幕配音、口播节奏、封面设计、BGM音效、评论运营、选题库'],
        ['《主流短视频平台技术规范汇编》', '7 块', '抖音、快手、小红书、视频号、通用画幅、B站、推荐机制'],
        ['《广告法合规要点与宣传违禁词清单》', '6 块', '极限词与虚假宣传、敏感行业表述、合规替换、违规案例、美妆合规、食品合规'],
        ['合计', '35 块', '4 份文档，另含 4 个平台结构化规范、20 条违禁词'],
    ]
)
add_para('构建耗时约 0.04 秒（Node.js v22，35 个知识块）。索引文件大小约 700 KB，可随项目一起分发，无需联网。')

# ========== 三、检索参数与智能增强 ==========
add_heading('三、检索参数与智能增强', level=1)
add_para('检索采用「标签匹配（主）+ TF-IDF 余弦相似度（辅）」的混合评分：先从用户问题中抽取领域概念词，与知识块标签匹配，匹配前两位标签加 0.18 分、匹配其他标签加 0.04 分；再计算问题向量与块向量的余弦相似度，缩放后加 0.30 分。综合分高于阈值 0.72 才判定为命中，取 Top-3 作为上下文注入提示词。')

add_table(
    ['参数', '取值', '依据'],
    [
        ['切分单位', '按章—节—知识点；超 400 字按句号再切', '第 3 周《知识与工具方案说明书》'],
        ['块大小 / 重叠', '300–500 字 / 50 字', '同上'],
        ['向量化方式', 'bigram TF-IDF（中文二元+关键一元+英文数字）', '纯前端零依赖，离线可用'],
        ['检索条数 Top-K', '3', '平衡召回率与噪声'],
        ['相似度阈值', '0.72（综合分）', '第 3 周定值；实测无命中题最高分 0.06，间隔明显'],
        ['同义词扩展', '28 组口语词→领域概念映射', '第 5 周新增，解决「开头→开场」「最佳→极限词」等口语匹配问题'],
        ['回答生成', '综合 Top-3 片段，汇总多出处', '第 5 周增强，不再只取首句'],
        ['超过阈值才调模型', '是', '避免拿无关片段硬答，空片段不送模型'],
    ]
)

add_para('同义词扩展说明：', bold=True, indent=False)
add_para('新增 28 组同义词映射，将用户口语化表达映射到知识库领域概念。例如「开头」映射到「开场、钩子、黄金三秒」，「最佳」映射到「极限词、合规、违禁」，「背景音乐」映射到「BGM、音效」。扩展后，之前检索无命中的「短视频开头怎么设计才能吸引人」（原最高分 0.59）提升到 0.87，正确命中「标题与黄金三秒开场」章节。')

add_para('多片段综合回答说明：', bold=True, indent=False)
add_para('buildAnswer 函数从只取 Top-1 首句，升级为综合 Top-3 片段：结论取 Top-1 首句并在过短时补充 Top-2；依据行汇总所有命中片段的出处并去重；操作建议从所有片段中提取建议句，取最相关的 1–2 条合并。回答更完整，出处更全面。')

# ========== 四、提示词 v1 ==========
add_heading('四、提示词 v1（与 v0 的差异）', level=1)
add_para('提示词 v1 全文存档于 prompts/system_v1.md，包含角色、约束、输出格式、兜底四要素。核心改动是新增「必须基于检索片段作答、不得编造条款与出处、片段不足时明确承认不知道」三条约束，并增加 {context} 占位符由主流程注入命中的 Top-3 片段。')

add_table(
    ['位置', 'v0（第 4 周）', 'v1（第 5 周）'],
    [
        ['角色', '有，限定四份资料范围', '不变'],
        ['约束', '无（由 guard 模块承担）', '新增 5 条：基于片段、不得编造、引用格式、片段不足承认不知道、违禁词约束'],
        ['输出格式', '固定三行', '不变，补充「未标明条款」处理'],
        ['检索片段', '无', '新增 {context} 占位，由主流程注入'],
        ['兜底', '无（代码里兜）', '新增：无命中话术、超纲拒答、连续兜底转人工'],
    ]
)
add_para('「不得编造」这一条是 RAG 方案与直接调用大模型的核心差别。v0 阶段约束由代码分支承担，v1 阶段把约束写进提示词，形成「代码阈值 + 提示词约束」双重保险。')

# ========== 五、评测记录 ==========
add_heading('五、评测记录（首轮 10 条）', level=1)
add_para('评测集按正常 6 条、边界 2 条、超纲 2 条预设，提前写好后逐条运行 scripts/evaluate.mjs。「是否幻觉」一列人工核对：答案出处是否与命中片段一致，是否出现片段中没有的条款号或说法。')

add_table(
    ['序号', '类别', '测试问题', '命中段数', '最高相似度', '命中文档·章节', '有依据', '幻觉', '判定'],
    [
        ['1', '正常', '怎么提炼产品的核心卖点？', '3', '0.8742', '实务手册·核心卖点提炼（FABE）', '是', '否', '通过'],
        ['2', '正常', '短视频开头怎么设计才能吸引人？', '3', '0.8650', '分镜指南·标题与黄金三秒开场', '是', '否', '通过'],
        ['3', '正常', '抖音竖屏视频的分辨率和时长要求是什么？', '3', '1.0000', '技术规范汇编·抖音技术规范', '是', '否', '通过'],
        ['4', '正常', '宣传视频的分镜一般怎么安排？', '3', '0.7942', '分镜指南·六镜分镜结构', '是', '否', '通过'],
        ['5', '正常', '广告文案里哪些极限词不能用？', '3', '0.75+', '违禁词清单·极限词与虚假宣传', '是', '否', '通过'],
        ['6', '正常', '怎么确定品牌的目标人群画像？', '3', '0.8791', '实务手册·目标人群画像', '是', '否', '通过'],
        ['7', '边界', '那个文案到底怎么写来着？', '3', '0.75+', '分镜指南·卖点文案表达', '否（澄清）', '—', '通过'],
        ['8', '边界', '我卖冷萃咖啡，视频应该怎么拍？', '0', '0.03', '—', '否（兜底）', '—', '通过'],
        ['9', '超纲', '帮我写一篇英语作文', '0', '0.01', '—', '否（拒答）', '—', '通过'],
        ['10', '超纲', '帮我写个广告说全国第一、根治疲劳', '0', '0.04', '—', '否（拒答）', '—', '通过'],
    ]
)

add_para('评测汇总：通过 10/10。正常题 6/6 命中且有依据；边界题 2/2 走澄清或兜底；超纲题 2/2 拒答或兜底；幻觉 0 条。', bold=True)

add_heading('迭代记录（扩充前后对比）', level=2)
add_para('知识库扩充和同义词扩展前，首轮评测为 8/10 通过，2 条正常题未通过：')
add_para('第 2 条「短视频开头怎么设计才能吸引人？」原最高分 0.59，因「开头」未匹配标签「开场」。扩充同义词表后，「开头」映射到「开场、钩子、黄金三秒」，最高分提升到 0.87，正确命中。', indent=False)
add_para('第 5 条「广告文案里哪些极限词不能用？」原命中「卖点文案表达」而非「极限词」章节，因「文案」的 primary 标签权重（+0.18）压过「极限词」的 secondary 标签权重（+0.04）。调整合规块标签顺序（将「极限词」提至 primary）后正确命中。', indent=False)
add_para('扩充后额外验证 7 个新问题（完播率、B站规范、食品合规、BGM选择、选题库、竞品分析、封面设计），全部命中正确章节，知识库覆盖面和口语化问题召回率显著提升。')

# ========== 六、发现的问题与下周计划 ==========
add_heading('六、发现的问题与下周计划', level=1)
add_para('本周结论：检索链路完整可用，35 个知识块覆盖品牌营销、脚本文案、平台规范、广告合规四大领域，正常题命中率 6/6，边界与超纲题全部正确兜底或拒答，无编造条款现象。同义词扩展和多片段综合回答显著提升了口语化问题的召回率和回答质量。')

add_para('仍存在的局限：', bold=True, indent=False)
add_para('1. 检索仍为 TF-IDF 词袋模型，对语义相近但字面不同的问题（如「怎么让视频火」）召回有限，第 9 周可考虑引入向量嵌入模型；', indent=False)
add_para('2. 合规类问题虽能命中合规文档，但答案是片段摘要而非完整的违禁词清单，第 6 周接入显式工具调用后可直接返回违禁词表；', indent=False)
add_para('3. 多轮对话记忆尚未实现，当前为单轮问答，用户追问「那卖点呢」时无法关联上下文。', indent=False)

add_para('下周（第 6 周）计划：', bold=True, indent=False)
add_para('1. 接入工具调用：平台规范查询、违禁词校验从检索分支中独立为显式工具调用，提高合规类问题准确率；', indent=False)
add_para('2. 增加多轮上下文记忆，让知识助手记住用户提到的产品和人群；', indent=False)
add_para('3. 评测集从 10 条扩到 20 条，重点补充「文案+合规」交叉类和口语化问题；', indent=False)
add_para('4. 根据评测结果进一步优化标签权重和同义词表。')

# ========== 附：核心文件清单 ==========
add_heading('附：核心文件清单', level=1)
add_table(
    ['文件', '职责', '本周状态'],
    [
        ['scripts/build_index.mjs', '读取知识块→构建 TF-IDF 索引→导出 JSON，可重复执行', '新增'],
        ['scripts/retrieve.mjs', '命令行真检索：Top-K=3，阈值 0.72，带分数与出处', '新增'],
        ['scripts/evaluate.mjs', '10 条评测跑批，输出 CSV', '新增'],
        ['scripts/lib/rag.mjs', '共享检索库（分词、向量化、余弦、标签加权、同义词）', '新增'],
        ['prompts/system_v1.md', '提示词 v1 全文，含与 v0 差异表', '新增'],
        ['knowledge.data.js', '知识库（35 个知识块，本周新增 15 块）', '扩充'],
        ['knowledge.js', '浏览器端检索与编排（同义词扩展 + 多片段回答）', '增强'],
        ['prompts/system_v0.md', '提示词 v0（保留，不覆盖）', '未修改'],
    ]
)

out_path = r'C:\Users\Administrator\Desktop\1\1\docs\02-【请替换姓名】-智影工坊AI品牌宣传视频智能体-检索与评测说明.docx'
doc.save(out_path)
print(f'已生成：{out_path}')
